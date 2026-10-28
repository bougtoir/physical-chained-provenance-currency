"""Regression tests for the DOCX build pipeline: table-cell rich rendering,
OMML conversion, Markdown/code-span stripping, and layout (docGrid)."""
import re
import sys
import zipfile
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "analysis"))
import build_docx as b  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "TFSC_SUBMISSION_FINAL" / "TFSC_MAIN_MANUSCRIPT_FINAL.docx"

RAW_LATEX = [r"\\mathrm", r"\\frac", r"\\approx", r"\\rightarrow",
             r"\\\(", r"\\\)", r"\\\[", r"\\\]", r"_\{", r"\^\{"]


@pytest.fixture(scope="module")
def built():
    b.main()
    assert DOCX.exists()
    return zipfile.ZipFile(DOCX).read("word/document.xml").decode()


def _visible_text(xml):
    return "".join(re.findall(r"<(?:w|m):t[^>]*>([^<]*)</(?:w|m):t>", xml))


def test_no_markdown_artifacts(built):
    vis = _visible_text(built)
    assert "**" not in vis
    assert "`" not in vis
    # stray single asterisks outside math are not allowed; '*' inside OMML
    # (e.g. r_t^*) is fine, so only check w:t (Word text) segments
    wtext = "".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", built))
    assert "*" not in wtext
    assert "| F" not in wtext  # markdown table pipes


def test_no_raw_latex(built):
    for pat in RAW_LATEX:
        assert not re.findall(pat, built), pat


def test_math_is_omml(built):
    assert built.count("<m:oMath>") > 100
    assert built.count("<m:oMathPara") >= 5


def test_table_cells_render_rich(built):
    # Table 5 cells: '**active**' must render as bold 'active' run, not literals
    from docx import Document
    doc = Document(DOCX)
    found = {}
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                txt = cell.text
                if "active" in txt and "low-separability" in txt:
                    runs = [r for p in cell.paragraphs for r in p.runs]
                    found["active"] = any(r.text == "active" and r.font.bold
                                          for r in runs)
                if "confirmed empirically" in txt:
                    runs = [r for p in cell.paragraphs for r in p.runs]
                    found["confirmed"] = any(
                        r.text == "confirmed empirically" and r.font.bold
                        for r in runs)
    assert found.get("active"), "bold 'active' run missing in Table 5"
    assert found.get("confirmed"), "bold 'confirmed empirically' run missing"


def test_table_cell_math_is_omml(built):
    # CRT_{0.5}=0 in Table 5 must be OMML inside a table cell
    from docx import Document
    doc = Document(DOCX)
    M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath"
    cells_with_math = [
        cell for t in doc.tables for row in t.rows for cell in row.cells
        if cell._element.findall(f".//{M}")]
    # F1/F2/F5/F6 rows of Table 5 all carry inline math
    assert len(cells_with_math) >= 4, "OMML missing inside table cells"


def test_no_docgrid(built):
    assert "docGrid" not in built
    assert "linePitch" not in built
