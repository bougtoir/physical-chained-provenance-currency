#!/usr/bin/env python3
"""Generate cover letter (md + docx) and highlights from canonical results.

All analysis-derived numbers are injected from
results/canonical_results.json — no retyped literals. The cover letter
keeps the approved two-page layout: a page break before the TFSC-fit
paragraph; no docGrid.
"""
import json
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
CANON = ROOT / "results" / "canonical_results.json"
OUTDIR = ROOT / "TFSC_SUBMISSION_FINAL"
FONT = "Times New Roman"
TOKEN_RE = re.compile(r"\{\{([A-Za-z0-9_]+)\}\}")

COVER_MD = """Dear Editors,

I submit "When Authentication Technologies Expire: Technology-Dependent Institutional Transitions in a Physical-Provenance Currency" for *Technological Forecasting and Social Change*. Institutions that rely on costly physical imitation — banknote printing, provenance certification, tamper evidence — face a moving technological target: an authentication architecture that is credible today may become economically obsolete as assay and fabrication technologies improve. The existing literature largely asks how secure a given institution is at a given technological state. This manuscript instead develops a technology-dependent institutional-transition model, asking which authentication institution remains viable as those costs move, and when an issuer should adapt, switch, or exit.

I instantiate and calibrate the model as a Physical Chained Provenance Currency — an extreme test case in which an issuer batches an abundant substrate with a public, dated, isotopically fingerprinted mixture — using {{n_samples}} measured iron-meteorite fingerprints, Monte-Carlo and closed-form analysis, four technology trajectories, and endogenous issuer adaptation under a switching-cost objective.

The strongest results are structural. Under symmetric technological progress the optimal sequence is fixed → rotating → exit, and adaptive switching extends viability by {{viability_ext_T1_pct}}% of the modeled horizon. Lineage-preserving authentication is optimal only in a narrow {{T2_r3_steps}}-step window — a delaying corridor, not a destination. Bulk predecessor-metal chaining adds no meaningful security once adversaries can reproduce the final bulk state: physical inheritance is not automatically institutional memory. And sufficiently cheap synthetic fabrication produces an ultimate expiration frontier near the model-specific c_syn/V ≈ {{expiration_frontier}}. Two further findings are honest negatives: adaptation is dominated by earlier exit under attacker-favoring progress, and neither hysteresis nor path dependence arises — transitions are viability-forced.

This belongs in TFSC because the contribution is not a speculative monetary design but a formal model of technology-driven institutional transition: viability frontiers are cost-ratio surfaces, regime changes are endogenous, and adaptation has finite benefits — a structure that applies to any institution whose credibility rests on the relative costs of authorized update, adversarial reconstruction, and verification.

The manuscript is original and not under consideration elsewhere; all code and reproduction materials accompany the submission. During preparation I used Devin (Cognition AI) to assist with simulation-code development and drafting; I reviewed and take full responsibility for all content.

Sincerely,

Tatsuki Onishi
Shiga University, Data Science and AI Innovation Research Promotion Center (DSAI)
"""

HIGHLIGHTS = """- Technological progress changes the optimal authentication institution.
- Adaptive switching extends viability by {{viability_ext_T1_pct}}% under symmetric progress.
- Bulk predecessor-metal chaining adds no security against synthesis.
- Physical lineage is optimal only in a narrow technological window.
- Cheap fabrication creates an expiration frontier for provenance regimes.
"""


def canon():
    return {k: v["display"] for k, v in
            json.loads(CANON.read_text()).items()}


def inject(text, c):
    return TOKEN_RE.sub(lambda m: c[m.group(1)], text)


def set_run(r, size=12, italic=False):
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.italic = italic
    rpr = r._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rf.set(qn(a), FONT)


def build_docx(md):
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = FONT
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn('w:ascii'), FONT)
    st.element.rPr.rFonts.set(qn('w:hAnsi'), FONT)
    paras = md.strip().split("\n\n")
    for i, para in enumerate(paras):
        p = doc.add_paragraph()
        if para.startswith("This belongs in TFSC"):
            # page break before the TFSC-fit paragraph (2-page layout)
            br = OxmlElement('w:br')
            br.set(qn('w:type'), 'page')
            p._element.insert(0, OxmlElement('w:r'))
            p._element[0].append(br)
        for part in re.split(r'(\*[^*]+\*)', para):
            if part.startswith('*') and part.endswith('*'):
                set_run(p.add_run(part[1:-1]), italic=True)
            elif part:
                set_run(p.add_run(part))
    for sect in doc.sections:
        for dg in sect._sectPr.findall(qn('w:docGrid')):
            sect._sectPr.remove(dg)
    return doc


def main():
    c = canon()
    md = inject(COVER_MD, c)
    hl = inject(HIGHLIGHTS, c)
    OUTDIR.mkdir(exist_ok=True)
    (OUTDIR / "TFSC_COVER_LETTER_FINAL.md").write_text(md)
    (OUTDIR / "TFSC_HIGHLIGHTS.txt").write_text(hl)
    build_docx(md).save(OUTDIR / "TFSC_COVER_LETTER_FINAL.docx")
    print("wrote cover letter md/docx + highlights")


if __name__ == "__main__":
    main()
