#!/usr/bin/env python3
"""Build TFSC_MAIN_MANUSCRIPT_FINAL.docx from manuscript/manuscript.md.

Times New Roman 12 pt body / 10 pt captions; all math -> OMML via
latex2mathml + docx-equation. Figures embedded inline right after first
citation; tables 1-4, 6 rendered from tables/*.csv at first citation;
table 5 is the in-text falsification table.
"""
import csv
import json
import re
from pathlib import Path

import latex2mathml.converter
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx_equation import mathml_to_omml

ROOT = Path(__file__).resolve().parents[1]
MD = ROOT / "manuscript" / "manuscript.md"
FIG = ROOT / "figures"
TBL = ROOT / "tables"
OUT = ROOT / "TFSC_SUBMISSION_FINAL" / "TFSC_MAIN_MANUSCRIPT_FINAL.docx"
CANON = ROOT / "results" / "canonical_results.json"

FONT = "Times New Roman"

TOKEN_RE = re.compile(r'\{\{([A-Za-z0-9_]+)\}\}')


def load_canonical():
    return {k: v["display"] for k, v in
            json.loads(CANON.read_text()).items()}


def inject(text, canon):
    def sub(m):
        key = m.group(1)
        if key not in canon:
            raise KeyError(f"canonical key missing: {key}")
        return canon[key]
    return TOKEN_RE.sub(sub, text)

FIG_CAPTIONS = {
    1: "Figure 1. PCPC architecture: the issuer batches an abundant substrate with a public, dated, isotopically fingerprinted source mixture; optional predecessor inheritance closes the material chain.",
    2: "Figure 2. Attacker geometry in fingerprint space: the forgery is constrained to the attainable-source subset; source separation Δ and within-source scatter σ set the verification problem.",
    3: "Figure 3. Static regime comparison: Material Rotation Asymmetry (MRA) and counterfeit viability as functions of normalized synthesis cost c_syn/V for architectures M0-M3b.",
    4: "Figure 4. Forgery pass probability P_F as a function of partial source access q and effective dimensionality d_eff; verification binds only above Δ/σ ≈ {{verif_sep_bound}}.",
    5: "Figure 5. Inclusion-chaining net benefit (lineage floor minus verification burden and issuer cost) across inclusion fraction α and verification-burden coefficient; positive only in the low-burden corner.",
    6: "Figure 6. Phase diagram in (Δ/σ, c_syn/V): verification-limited, viable, economically forgeable, and technologically obsolete regions; the expiration frontier lies near c_syn/V ≈ {{expiration_frontier}}.",
    7: "Figure 7. Retrospective authentication: log10 Bayes factor for a forged chain link versus depth k, for coarse and near-perfect fabricators.",
    8: "Figure 8. Endogenous institutional transitions: optimal authentication regime over normalized technology-progression steps for trajectories T1-T4 (T1 symmetric; T2 attacker-favoring; T3 defender-favoring; T4 discontinuous breakthrough). Steps index technology states, not calendar time.",
    9: "Figure 9. Regime map: optimal regime in (c_syn/V, c_ver/V) space for three microfabrication-cost tiers, with T1-T4 trajectories overlaid.",
    10: "Figure 10. Value of adaptive switching versus the best fixed regime and versus immediate exit, per trajectory and switching-weight κ.",
}

TABLE_FILES = {
    1: ("table_regimes.csv", "Table 1. Static regime comparison at the calibrated synthesis-cost point: architecture, normalized synthesis cost, forgery pass probability, attacker and issuer cost ratios, and Material Rotation Asymmetry (MRA)."),
    2: ("table_crt05.csv", "Table 2. Coalition-size reversal: CRT_0.5 (fraction of runs where the mean/median forgery passes verification at least half the time) by coalition size N and source separation Δ/σ."),
    3: ("table_empirical_meta.csv", "Table 3. Empirical calibration summary: iron-meteorite fingerprint corpus (IEDA/AstroMat dataset series)."),
    4: ("table_transitions.csv", "Table 4. Endogenous transition summary per technology trajectory (issuer rule A, κ = 0): endogenous regime sequence, first transition step, number of transitions, R3-optimal window, adaptation value, and viability extension."),
    6: ("table_blockchain_comparison.csv", "Table 6. PCPC versus cryptographic blockchain: chronological linkage and tamper evidence are shared; consensus, replication and append-only guarantees are not claimed."),
}


def omml_runs(latex, paragraph, display=False):
    mathml = latex2mathml.converter.convert(latex)
    omml = mathml_to_omml(mathml)
    el = paragraph._element
    if display:
        from docx.oxml import OxmlElement
        mpara = OxmlElement('m:oMathPara')
        mpara.append(omml)
        el.append(mpara)
    else:
        el.append(omml)


def set_run(r, size=12, bold=False, italic=False):
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    rpr = r._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        from docx.oxml import OxmlElement
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:ascii'), FONT)
    rfonts.set(qn('w:hAnsi'), FONT)
    rfonts.set(qn('w:cs'), FONT)


MATH_RE = re.compile(r'(\\\(|\\\)|\\\[|\\\])')
EMPH_RE = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*)')
CODE_RE = re.compile(r'`([^`]+)`')


def add_rich(paragraph, text, size=12, base_bold=False):
    """Add text with \( \) inline math as OMML and *emphasis* runs."""
    pos = 0
    for m in MATH_RE.finditer(text):
        chunk = text[pos:m.start()]
        if chunk:
            _emph(paragraph, chunk, size, base_bold)
        if m.group(1) == '\\(':
            end = text.find('\\)', m.end())
            omml_runs(text[m.end():end].strip(), paragraph)
            pos = end + 2
        else:
            pos = m.end()
    _emph(paragraph, text[pos:], size, base_bold)


def set_code_run(r, size=11):
    r.font.name = 'Consolas'
    r.font.size = Pt(size)
    rpr = r._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        from docx.oxml import OxmlElement
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rfonts.set(qn(attr), 'Consolas')


def _emph(paragraph, chunk, size, base_bold):
    pos = 0
    for cm in CODE_RE.finditer(chunk):
        _emph_plain(paragraph, chunk[pos:cm.start()], size, base_bold)
        set_code_run(paragraph.add_run(cm.group(1)), size - 1)
        pos = cm.end()
    _emph_plain(paragraph, chunk[pos:], size, base_bold)


def _emph_plain(paragraph, chunk, size, base_bold):
    for part in EMPH_RE.split(chunk):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            set_run(paragraph.add_run(part[2:-2]), size, bold=True)
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            set_run(paragraph.add_run(part[1:-1]), size, italic=True)
        else:
            set_run(paragraph.add_run(part), size, bold=base_bold)


def add_table(doc, rows, header=True):
    t = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
    t.style = 'Table Grid'
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            if j >= len(t.rows[i].cells):
                continue
            p = t.rows[i].cells[j].paragraphs[0]
            add_rich(p, str(cell), size=10, base_bold=(header and i == 0))
    return t


def fmt_csv_table(path):
    rows = list(csv.reader(open(path)))
    if path.name == 'table_empirical_meta.csv':
        return [['Metric', 'Value']] + [[r[0], f"{float(r[1]):.4g}"] for r in rows[1:] if len(r) > 1]
    if path.name == 'table_crt05.csv':
        rows[0][0] = 'N / Δ/σ'
    if path.name == 'table_regimes.csv':
        return [[r[0], r[1][:6], r[2][:6], r[3][:6], r[4][:6], r[5][:6]] for r in rows]
    return rows


def main():
    canon = load_canonical()
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = FONT
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn('w:ascii'), FONT)
    st.element.rPr.rFonts.set(qn('w:hAnsi'), FONT)

    text = inject(MD.read_text(), load_canonical())
    # strip working-draft note
    text = re.sub(r'\*\(Working draft[^)]*\)\*\n\n', '', text)

    fig_done, tbl_done = set(), set()
    lines = text.split('\n')
    i = 0
    table_buf = None

    while i < len(lines):
        ln = lines[i]
        if ln.startswith('|') and '|' in ln[1:]:
            # markdown table block
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                if not re.match(r'^\|[\s\-|]+\|$', lines[i]):
                    rows.append([c.strip() for c in lines[i].strip('|').split('|')])
                i += 1
            add_table(doc, rows)
            continue
        if ln.startswith('# '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run(p.add_run(ln[2:].strip()), 16, bold=True)
            # author line
            p2 = doc.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run(p2.add_run('Tatsuki Onishi'), 12)
            p3 = doc.add_paragraph()
            p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run(p3.add_run('Shiga University, Data Science and AI Innovation Research Promotion Center (DSAI)'), 12, italic=True)
        elif ln.startswith('## '):
            p = doc.add_paragraph()
            set_run(p.add_run(ln[3:].strip()), 14, bold=True)
        elif ln.startswith('### '):
            p = doc.add_paragraph()
            set_run(p.add_run(ln[4:].strip()), 12, bold=True)
        elif re.match(r'^\d+\.\s', ln):
            p = doc.add_paragraph()
            add_rich(p, re.sub(r'^\d+\.\s*', '', ln))
        elif ln.strip() == '':
            pass
        elif ln.startswith('\\['):
            block = []
            i += 1
            while i < len(lines):
                if lines[i].startswith('\\]'):
                    break
                block.append(lines[i].strip())
                i += 1
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            latex = ' '.join(block).strip()
            try:
                omml_runs(latex, p, display=True)
            except Exception as e:
                set_run(p.add_run(latex), 12, italic=True)
                print('math fail:', latex[:60], e)
        elif ln.startswith('**') and ln.rstrip().endswith('**') and ln.count('**') == 2:
            # caption line inside body (e.g. Table 5 caption)
            p = doc.add_paragraph()
            add_rich(p, ln.replace('**', ''), size=10, base_bold=False)
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(10)
        else:
            p = doc.add_paragraph()
            add_rich(p, ln)
            # figure/table insertion after first citation
            for n in range(1, 11):
                if n not in fig_done and re.search(rf'\bFig(?:ure)?\.\s*{n}\b|\bFigure\s*{n}\b', ln):
                    files = list(FIG.glob(f'fig{n}_*.png'))
                    if files:
                        doc.add_picture(str(files[0]), width=Inches(5.8))
                        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                        cap = doc.add_paragraph()
                        add_rich(cap, inject(FIG_CAPTIONS[n], canon), size=10)
                        fig_done.add(n)
            for n in [1, 2, 3, 4, 6]:
                if n not in tbl_done and re.search(rf'\bTable\s*{n}\b', ln):
                    fn, cap = TABLE_FILES[n]
                    add_table(doc, fmt_csv_table(TBL / fn))
                    cap_p = doc.add_paragraph()
                    add_rich(cap_p, inject(cap, canon), size=10)
                    tbl_done.add(n)
        i += 1

    # remove East-Asian character grid (docGrid/linePitch) — TFSC uses
    # ordinary paragraph line spacing
    for sect in doc.sections:
        sectPr = sect._sectPr
        for dg in sectPr.findall(qn('w:docGrid')):
            sectPr.remove(dg)

    OUT.parent.mkdir(exist_ok=True)
    doc.save(OUT)
    print('figures inserted:', sorted(fig_done))
    print('tables inserted:', sorted(tbl_done))
    print('wrote', OUT)


if __name__ == '__main__':
    main()
